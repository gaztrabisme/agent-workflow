#!/usr/bin/env python3
"""
check_quoted_figures.py -- check that every figure quoted in an email, cover note
or slide exists in the file(s) it cites, before it goes out.

Why this exists: a figure retyped into a message drifts from the file it came
from, and nobody re-opens the attachment to compare. One email body priced three
environments at USD 10,156 while the attached workbook priced one at USD 3,157.15.
Both looked plausible; only a side-by-side shows the gap. This script does the
side-by-side for every number in the message.

What it does: extracts every number of at least --min from the text (currency
prefixes $, USD, RM, MYR; thousands separators; a k suffix, 60K = 60000), skipping
percentages, dates such as 13/10/2026, and bare years 1990-2100. It then looks for
each one among the values of the source files: cached numeric cell values of
.xlsx workbooks (not numbers inside text cells, such as "64 CU x 730 hours",
which matched by coincidence in testing), all text of .pptx decks, and
lines of .md/.txt files. A quoted figure matches a source value if the source
value, rounded to the precision the figure is displayed at, equals it (3,136
matches 3135.6; 3,157.15 matches 3157.1523; 60K matches 59,800), or if the two are
within --tolerance-pct percent. Signs are ignored: (1,234) and -1,234 match 1234.

What it cannot do: it matches values, not meaning. A figure can exist in the file
for a different period, environment or scope -- 3,136 can be one month in the file
and a whole phase in the email -- and the script will call it a MATCH. A person
reads every match against the scope the message cites, and every NO MATCH. Small
round numbers match by coincidence easily; raise --min to cut that noise. A
workbook saved by a tool that never calculated it has formulas with no cached
values; the script warns per sheet, because nothing in those cells can match.
A bare integer between 1990 and 2100 is treated as a year and skipped, even when
it is an amount.

Usage:
    python check_quoted_figures.py email.txt book.xlsx deck.pptx
    python check_quoted_figures.py note.md book.xlsx --tolerance-pct 0 --min 1000
    python check_quoted_figures.py slide.pptx book.xlsx --include-pct

    <text-file>  .txt .md .eml .docx .pptx  (the message being checked)
    <source>     .xlsx .pptx .md .txt        (the files it cites)

Exit code is 1 if any quoted figure has no match, 2 on a usage or read error,
else 0.
"""

import argparse
import email
import os
import re
import sys
import zipfile
from email import policy
from html import unescape

# --------------------------------------------------------------------------
# Number extraction
# --------------------------------------------------------------------------

INVISIBLE = dict.fromkeys(map(ord, "​‌‍⁠﻿"), None)
SPACES = {ord(c): " " for c in "    "}
MINUS = {ord("−"): "-"}

DATE_RE = re.compile(
    r"(?<!\d)(?:\d{1,2}[/.\-]\d{1,2}[/.\-](?:\d{4}|\d{2})|\d{4}-\d{1,2}-\d{1,2})(?!\d)")

NUM_RE = re.compile(
    r"(?<![\w.,/])"
    r"(?P<cur>(?:USD|US\$|MYR|RM|\$)\s?)?"
    r"(?P<num>\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)"
    r"(?P<suf>[kK](?![A-Za-z]))?"
    r"(?P<pct>\s?%)?"
    r"(?![\d,]\d)")


def clean(text):
    return text.translate(INVISIBLE).translate(SPACES).translate(MINUS)


def parse_numbers(text):
    """Yield dicts: value, unit (half of it is the rounding window), raw, flags."""
    text = DATE_RE.sub(lambda m: " " * len(m.group(0)), clean(text))
    for m in NUM_RE.finditer(text):
        num = m.group("num")
        mult = 1000.0 if m.group("suf") else 1.0
        decimals = len(num.split(".")[1]) if "." in num else 0
        value = float(num.replace(",", "")) * mult
        yield {
            "value": value,
            "unit": (10.0 ** -decimals) * mult,
            "raw": m.group(0).strip(),
            "pct": bool(m.group("pct")),
            "year": (not m.group("cur") and not m.group("suf") and "," not in num
                     and "." not in num and 1990 <= value <= 2100),
        }


# --------------------------------------------------------------------------
# Readers: each returns a list of (location, line_text)
# --------------------------------------------------------------------------


def _lines(text, prefix="line"):
    return [(f"{prefix} {i}", ln) for i, ln in enumerate(text.splitlines(), 1) if ln.strip()]


def read_plain(path):
    with open(path, encoding="utf-8", errors="replace") as fh:
        return _lines(fh.read())


def read_eml(path):
    with open(path, "rb") as fh:
        msg = email.message_from_binary_file(fh, policy=policy.default)
    part = msg.get_body(preferencelist=("plain", "html"))
    text = part.get_content() if part else ""
    if part is not None and part.get_content_type() == "text/html":
        text = unescape(re.sub(r"<br\s*/?>|</p>|</tr>|</td>|</div>", "\n", text, flags=re.I))
        text = re.sub(r"<[^>]+>", " ", text)
    return _lines(text)


def read_docx(path):
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    out = []
    for i, p in enumerate(re.findall(r"<w:p[ >].*?</w:p>", xml, flags=re.S), 1):
        t = unescape("".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", p)))
        if t.strip():
            out.append((f"paragraph {i}", t))
    return out


def _shape_texts(shape):
    if shape.shape_type == 6:  # group
        for s in shape.shapes:
            yield from _shape_texts(s)
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            yield "".join(r.text for r in p.runs)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            yield " | ".join(c.text for c in row.cells)


def read_pptx(path):
    from pptx import Presentation
    out = []
    for n, slide in enumerate(Presentation(path).slides, 1):
        for shape in slide.shapes:
            for t in _shape_texts(shape):
                if t and t.strip():
                    out.append((f"slide {n}", t))
    return out


def read_text_file(path):
    ext = os.path.splitext(path)[1].lower()
    return {".pptx": read_pptx, ".docx": read_docx, ".eml": read_eml}.get(ext, read_plain)(path)


def read_xlsx(path, warnings):
    """Return list of (location, value). Cached values only."""
    import openpyxl
    name = os.path.basename(path)
    vals = openpyxl.load_workbook(path, data_only=True)
    forms = openpyxl.load_workbook(path, data_only=False)
    out = []
    for ws in vals.worksheets:
        fws = forms[ws.title]
        uncached = 0
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                f = fws[c.coordinate].value
                if v is None:
                    if isinstance(f, str) and f.startswith("="):
                        uncached += 1
                    continue
                loc = f"{name} {ws.title}!{c.coordinate}"
                if not (isinstance(f, str) and f.startswith("=")):
                    loc += " [typed value, no formula]"
                if isinstance(v, bool):
                    continue
                if isinstance(v, (int, float)):
                    out.append((loc, float(v)))
        if uncached:
            warnings.append(f"WARNING {name} sheet '{ws.title}': {uncached} formula cell(s) "
                            f"have no cached value; nothing in them can match "
                            f"(open and save the file in Excel or LibreOffice).")
    return out


def read_source(path, warnings):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(path, warnings)
    name = os.path.basename(path)
    lines = read_pptx(path) if ext == ".pptx" else read_plain(path)
    return [(f"{name} {loc}", n["value"]) for loc, t in lines for n in parse_numbers(t)]


# --------------------------------------------------------------------------
# Matching
# --------------------------------------------------------------------------


def find_match(q, sources, tol_pct):
    """Closest source value: within rounding of the quoted precision ("exact"), else within tol_pct."""
    qv = abs(q["value"])
    window = q["unit"] / 2 + 1e-9
    best = near = None
    for loc, sv in sources:
        sv = abs(sv)
        d = abs(sv - qv)
        if d <= window:
            if best is None or d < best[0]:
                best = (d, loc, sv)
        elif tol_pct > 0 and qv and d / qv * 100 <= tol_pct:
            if near is None or d < near[0]:
                near = (d, loc, sv)
    if best:
        return best[1], best[2], "exact"
    if near:
        return near[1], near[2], f"within {near[0] / qv * 100:.2f}%"
    return None


def fmt(v):
    return f"{v:,.4f}".rstrip("0").rstrip(".")


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[1],
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("text_file")
    ap.add_argument("sources", nargs="+")
    ap.add_argument("--tolerance-pct", type=float, default=0.5)
    ap.add_argument("--min", type=float, default=100.0)
    ap.add_argument("--include-pct", action="store_true", help="also check percentages")
    a = ap.parse_args()

    warnings = []
    try:
        lines = read_text_file(a.text_file)
        sources = []
        for s in a.sources:
            sources.extend(read_source(s, warnings))
    except Exception as e:  # noqa: BLE001
        print(f"ERROR reading input: {e}")
        return 2

    for w in warnings:
        print(w)
    print(f"Sources: {len(a.sources)} file(s), {len(sources)} numeric value(s). "
          f"Tolerance {a.tolerance_pct}%, min {fmt(a.min)}.\n")

    matched = unmatched = 0
    for loc, text in lines:
        line = " ".join(clean(text).split())
        if len(line) > 90:
            line = line[:87] + "..."
        for q in parse_numbers(text):
            if q["year"] or (q["pct"] and not a.include_pct) or abs(q["value"]) < a.min:
                continue
            hit = find_match(q, sources, a.tolerance_pct)
            if hit:
                matched += 1
                hloc, hv, how = hit
                res = f"MATCH     {hloc} = {fmt(hv)}" + ("" if how == "exact" else f" ({how})")
            else:
                unmatched += 1
                res = "NO MATCH"
            print(f"{q['raw']:>16}  {res}\n{'':>16}  [{loc}] {line}")

    print(f"\nSummary: {matched} matched, {unmatched} unmatched, "
          f"{matched + unmatched} quoted figure(s) checked.")
    if unmatched:
        print("Read every NO MATCH. Read every MATCH against the scope the message cites.")
        print("A match on a [typed value, no formula] cell is a figure nothing computes: check it against the cells that should.")
    return 1 if unmatched else 0


if __name__ == "__main__":
    sys.exit(main())
