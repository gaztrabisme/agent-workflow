#!/usr/bin/env python3
"""Negative tests for build.py and verify.py, on the fixture in tests/fixtures with project.yaml.example.

Part 0: the untouched fixture builds, recalculates and verifies (exit 0).
Part 1: each broken input (data or settings) must be refused by the build: exit 2 and the expected message.
Part 2: each input that builds but breaks a rule: the verifier must exit 1 and name the expected group.
Part 3: each workbook edited after the build (then recalculated): the verifier must exit 1 and name the group.

    python tests/run_negative.py

Exit 0 only if every case behaves as expected. Work happens in a temporary folder; nothing here changes.
Needs LibreOffice (found as in wbs_common.find_soffice).
"""
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl
import yaml
from openpyxl.styles import PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FIXTURE = os.path.join(HERE, "fixtures")
EXAMPLE = os.path.join(ROOT, "project.yaml.example")
BUILD = os.path.join(ROOT, "build.py")
VERIFY = os.path.join(ROOT, "verify.py")
sys.path.insert(0, ROOT)
import wbs_common  # noqa: E402

ENV = dict(os.environ, PYTHONUTF8="1")
PY = sys.executable
TEST_NAME = "Jordan"   # a person name the test settings forbid on shown sheets


def run(cmd):
    p = subprocess.run([PY, "-B"] + cmd, env=ENV, capture_output=True, text=True, encoding="utf-8",
                       errors="replace")
    return p.returncode, (p.stdout or "") + (p.stderr or "")


def load(path):
    with open(path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def dump(path, doc):
    with open(path, "w", encoding="utf-8") as f:
        yaml.safe_dump(doc, f, allow_unicode=True, sort_keys=False)


def edit(d, name, fn):
    path = os.path.join(d, name)
    doc = load(path)
    fn(doc)
    dump(path, doc)


def make_case(work, tag):
    """A fresh copy of the fixture data plus a test project.yaml inside it."""
    d = os.path.join(work, tag)
    shutil.copytree(FIXTURE, d)
    cfg = load(EXAMPLE)
    cfg["text"]["person_names"] = [TEST_NAME]
    dump(os.path.join(d, "project.yaml"), cfg)
    return d


def build(d, out):
    return run([BUILD, "--project", os.path.join(d, "project.yaml"), "--data", d, "--out", out])


def verify(xlsx, d, extra=()):
    return run([VERIFY, xlsx, "--project", os.path.join(d, "project.yaml"), "--data", d, "--max-lines", "30"]
               + list(extra))


# ---- Part 1: inputs the build must refuse (exit 2) ----------------------------------------------
def status_item(c, task_id):
    return next(s for s in c["status"] if s["task_id"] == task_id)


BUILD_CASES = [
    ("removal without a source", lambda d: edit(d, "changes.yaml", lambda c: status_item(c, "1.3.1.1").pop("source")),
     "has no source"),
    ("not-needed status without a source",
     lambda d: edit(d, "changes.yaml", lambda c: status_item(c, "1.4.1.1").pop("source")), "has no source"),
    ("status names an unknown task", lambda d: edit(d, "changes.yaml", lambda c: c["status"][0].update(task_id="9.9.9.9")),
     "unknown task id"),
    ("status has a bad value", lambda d: edit(d, "changes.yaml", lambda c: c["status"][0].update(status="maybe")),
     "bad status"),
    ("status has a bad owner", lambda d: edit(d, "changes.yaml", lambda c: c["status"][0].update(owner="Acme")),
     "bad owner"),
    ("status change on an overhead row", lambda d: edit(d, "changes.yaml", lambda c: c["status"].append(
        {"task_id": "1.6.1.1", "status": "in", "owner": "Us", "reason": "x"})), "overhead row"),
    ("added task points at an unknown WBS row",
     lambda d: edit(d, "changes.yaml", lambda c: c["add"][0].update(wbs_id="Z.9")), "unknown wbs_id"),
    ("added task without a source", lambda d: edit(d, "changes.yaml", lambda c: c["add"][0].pop("source")),
     "has no source"),
    ("added task id without its prefix", lambda d: edit(d, "changes.yaml", lambda c: c["add"][0].update(id="Y-01")),
     "does not start with"),
    ("proposed task has an unknown role", lambda d: edit(d, "proposals.yaml", lambda a: a["add"][0].update(role="Wizard")),
     "unknown role"),
    ("proposed task reuses an ID", lambda d: edit(d, "proposals.yaml", lambda a: a["add"][0].update(id="X-01")),
     "already used"),
    ("proposed change names an unknown task",
     lambda d: edit(d, "proposals.yaml", lambda a: a["change"][0].update(task_id="9.9.9.9")), "unknown task"),
    ("proposed change on an out-of-scope task",
     lambda d: edit(d, "proposals.yaml", lambda a: a["change"][0].update(task_id="1.3.1.1")), "cannot also be changed"),
    ("proposed effort is text", lambda d: edit(d, "proposals.yaml", lambda a: a["change"][0].update(effort="x")),
     "effort is not a number"),
    ("criteria file names an unknown task", lambda d: open(os.path.join(d, "task_criteria_9.yaml"), "w").write(
        "tasks:\n- task_id: 9.9.9.9\n  subtasks: []\n  acs: []\n"), "unknown task"),
    ("criteria for one task in two files", lambda d: shutil.copyfile(
        os.path.join(d, "task_criteria_1.yaml"), os.path.join(d, "task_criteria_2.yaml")), "listed twice across"),
    ("plan feature missing from the map", lambda d: edit(d, "feature_map.yaml", lambda m: m.update(
        map=[x for x in m["map"] if x["plan_feature"] != "1.1.1"])), "is not in feature_map.yaml"),
    ("story maps to an unknown WBS row", lambda d: edit(d, "stories.yaml", lambda s: s["stories"][0].update(wbs_ids=["Z.9"])),
     "maps to unknown WBS row"),
    ("dependency blocks an unknown WBS row", lambda d: edit(d, "dependencies.yaml", lambda s: s[0].update(blocks=["Z.9"])),
     "blocks unknown WBS row"),
    ("dependency from an unknown party", lambda d: edit(d, "dependencies.yaml", lambda s: s[0].update({"from": "Acme"})),
     "from must be one of"),
    ("changes file is not valid YAML", lambda d: open(os.path.join(d, "changes.yaml"), "w").write("status: [\n"),
     "not valid YAML"),
    ("plan tasks file is missing", lambda d: os.remove(os.path.join(d, "plan_tasks.yaml")), "missing input file"),
    ("settings ask for the pricing model", lambda d: edit(d, "project.yaml", lambda c: c.update(pricing="model")),
     "pricing: model is not in the first cut"),
    ("settings miss the plan total", lambda d: edit(d, "project.yaml", lambda c: c["plan"].pop("total")),
     "missing setting plan.total"),
    ("proposal fill equals the fill-in fill", lambda d: edit(d, "project.yaml", lambda c: c["marking"]["fills"].update(
        proposal=c["marking"]["fills"]["fill_in"])), "must differ"),
    ("settings miss a Task Detail column", lambda d: edit(d, "project.yaml", lambda c: c["columns"]["task_detail"].pop()),
     "columns.task_detail must list"),
]


# ---- Part 2: inputs that build, but the verifier must reject --------------------------------------
def sub_sum(t):
    t["tasks"][0]["subtasks"][1]["days"] = 1.0


def sub_over(t):
    t["tasks"][1]["subtasks"] = [{"name": "Build the service", "days": 3.5}, {"name": "Tests", "days": 1.0}]


VERIFY_CASES = [
    ("proposal reason without the prefix",
     lambda d: edit(d, "proposals.yaml", lambda a: a["change"][0].update(reason="fixture reason")), "proposals"),
    ("proposed effort breaks the effort rule",
     lambda d: edit(d, "proposals.yaml", lambda a: a["change"][0].update(effort=4.0)), "effort rule"),
    ("proposed change without a source",
     lambda d: edit(d, "proposals.yaml", lambda a: a["change"][0].pop("source")), "sources"),
    ("removal source without the row marker",
     lambda d: edit(d, "changes.yaml", lambda c: status_item(c, "1.3.1.1").update(source="Q&A log, line 22")),
     "removals"),
    ("sub-task days do not sum to effort", lambda d: edit(d, "task_criteria_1.yaml", sub_sum), "criteria coverage"),
    ("sub-task over the day limit", lambda d: edit(d, "task_criteria_1.yaml", sub_over), "criteria coverage"),
    ("task over the threshold has no criteria", lambda d: edit(d, "task_criteria_1.yaml", lambda t: t.update(
        tasks=[x for x in t["tasks"] if x["task_id"] != "1.1.1.1"])), "criteria coverage"),
    ("dependency has no if-late text", lambda d: edit(d, "dependencies.yaml", lambda s: s[0].update(if_late="")),
     "dependencies"),
    ("target without the test-method clause", lambda d: edit(d, "feature_map.yaml", lambda m: m["rows"][0].update(
        acceptance="Match rate at 95% or more.")), "acceptance clause"),
    ("WBS description over the cap", lambda d: edit(d, "feature_map.yaml", lambda m: m["rows"][0].update(
        description="x " * 90)), "text length"),
    ("added story criterion over the cap", lambda d: edit(d, "stories.yaml", lambda s: s["stories"][1].update(
        acs=["y" * 160])), "text length"),
    ("person name in a WBS note", lambda d: edit(d, "feature_map.yaml", lambda m: m["rows"][0].update(
        note="Checked by %s" % TEST_NAME)), "person names and sources"),
    ("internal file name in an assumption", lambda d: edit(d, "assumptions.yaml", lambda a: a[0].update(
        text="We assume one site only, as in scope_notes.yaml.")), "person names and sources"),
    ("owner's plan total not reproduced", lambda d: edit(d, "project.yaml", lambda c: c["plan"].update(total=30.0)),
     "reference check"),
    ("story given verbatim but changed", lambda d: edit(d, "stories.yaml", lambda s: s["stories"][0].update(
        acs=["Entrant is matched quickly", "Result is shown to the guard"])), "stories"),
]


# ---- Part 3: workbook tampering (edit, recalculate, verify) -------------------------------------
def find_row(ws, column, value):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, column).value == value:
            return r
    raise LookupError(value)


ORANGE = PatternFill("solid", start_color="FFF8CBAD", end_color="FFF8CBAD")


def t_name(wb):
    ws = wb["WBS"]
    ws.cell(find_row(ws, 2, "A.1"), 11).value = "Checked by %s" % TEST_NAME


def t_order(wb):
    wb._sheets = list(reversed(wb._sheets))


def t_factor(wb):
    ws = wb["Master"]
    ws.cell(9, 1).value = "Productivity factor"
    ws.cell(9, 2).value = 0.8


def t_viet(wb):
    wb["Task Detail"].cell(2, 6).value = "Kiểm tra"


def t_story(wb):
    wb["User Stories"].cell(2, 4).value = "1. Entrant is matched in 9 s or less."


def t_wbs_typed(wb):
    ws = wb["WBS"]
    ws.cell(find_row(ws, 2, "A.1"), 10).value = 99


def t_total_typed(wb):
    ws = wb["Effort Summary"]
    ws.cell(find_row(ws, 1, "TOTAL"), 2).value = 1


def t_error(wb):
    ws = wb["WBS"]
    ws.cell(find_row(ws, 2, "A.1"), 10).value = "=1/0"


def t_glossary(wb):
    ws = wb["References"]
    r = find_row(ws, 2, "Add-on")
    for c in range(1, 4):
        ws.cell(r, c).value = None


def t_dup(wb):
    ws = wb["Task Detail"]
    last = ws.max_row + 1
    for c in range(1, 18):
        ws.cell(last, c).value = ws.cell(2, c).value


def t_fill_removed(wb):
    ws = wb["Task Detail"]
    r = find_row(ws, 1, "1.1.1.2")
    for c in range(1, 18):
        ws.cell(r, c).fill = PatternFill(fill_type=None)


def t_fill_added(wb):
    ws = wb["Task Detail"]
    ws.cell(find_row(ws, 1, "1.1.1.3"), 11).fill = ORANGE


def t_legend(wb):
    ws = wb["References"]
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == "Colour legend":
            for c in range(1, 4):
                ws.cell(r, c).value = None


def t_figure(wb):
    ws = wb["Task Detail"]
    ws.cell(find_row(ws, 1, "1.1.1.3"), 11).value = 2.0


def t_prefix_note(wb):
    ws = wb["Task Detail"]
    r = find_row(ws, 1, "1.1.1.2")
    ws.cell(r, 17).value = "fixture reason without the prefix"


WB_CASES = [
    ("person name in a shown cell", t_name, "person names and sources"),
    ("sheets out of order", t_order, "structure"),
    ("forbidden factor on Master", t_factor, "master inputs"),
    ("Vietnamese letters in a task name", t_viet, "diacritics"),
    ("verbatim story criterion changed", t_story, "stories"),
    ("WBS sum typed over with a number", t_wbs_typed, "live formulas"),
    ("summary TOTAL typed over", t_total_typed, "live formulas"),
    ("formula returns an error value", t_error, "cached values"),
    ("task row duplicated", t_dup, "task detail rows"),
    ("glossary term removed from References", t_glossary, "references"),
    ("proposal fill removed from a proposed change", t_fill_removed, "proposal marking"),
    ("proposal fill on a row that is not a proposal", t_fill_added, "proposal marking"),
    ("colour legend removed", t_legend, "proposal marking"),
    ("task effort changed by hand", t_figure, "effort rule"),
    ("proposal note lost its prefix", t_prefix_note, "proposal marking"),
]


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    work = tempfile.mkdtemp(prefix="wbs_neg_")
    results = []
    try:
        base = make_case(work, "base")
        base_x = os.path.join(work, "base.xlsx")
        code, out = build(base, base_x)
        if code != 0:
            print("Part 0 FAIL: fixture build exit %d\n%s" % (code, out))
            return 1
        wbs_common.recalc(base_x)
        code, out = verify(base_x, base)
        results.append(("part 0: fixture verifies", code == 0, out.strip().splitlines()[-1] if out else ""))

        for i, (name, fn, expect) in enumerate(BUILD_CASES):
            d = make_case(work, "b%02d" % i)
            fn(d)
            code, out = build(d, os.path.join(work, "b%02d.xlsx" % i))
            ok = code == 2 and expect in out
            results.append(("build refuses: " + name, ok, "" if ok else "exit %d: %s" % (code, out.strip()[-160:])))

        for i, (name, fn, group) in enumerate(VERIFY_CASES):
            d = make_case(work, "v%02d" % i)
            fn(d)
            x = os.path.join(work, "v%02d.xlsx" % i)
            code, out = build(d, x)
            if code != 0:
                results.append(("verifier names: " + name, False, "build exit %d: %s" % (code, out.strip()[-160:])))
                continue
            wbs_common.recalc(x)
            code, out = verify(x, d)
            ok = code == 1 and ("FAIL %s (" % group) in out
            results.append(("verifier names: " + name, ok, "" if ok else "exit %d, expected FAIL %s" % (code, group)))

        for i, (name, fn, group) in enumerate(WB_CASES):
            x = os.path.join(work, "w%02d.xlsx" % i)
            wb = openpyxl.load_workbook(base_x)
            fn(wb)
            wb.save(x)
            wbs_common.recalc(x)
            code, out = verify(x, base)
            ok = code == 1 and ("FAIL %s (" % group) in out
            results.append(("tampered workbook: " + name, ok, "" if ok else "exit %d, expected FAIL %s" % (code, group)))

        code, out = verify(base_x, base, ["--plan-total", "33"])
        ok = code == 1 and "FAIL reference check (" in out
        results.append(("verifier names: wrong plan total on the command line", ok, "" if ok else "exit %d" % code))
    finally:
        shutil.rmtree(work, ignore_errors=True)

    failed = 0
    for name, ok, note in results:
        print("%s  %s%s" % ("ok  " if ok else "FAIL", name, "" if ok else "  (" + note + ")"))
        failed += 0 if ok else 1
    print("%d of %d cases behave as expected" % (len(results) - failed, len(results)))
    return 0 if failed == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
