#!/usr/bin/env python3
"""Point each requirement at the places in the proposal that most plausibly answer it.

From the technical-proposal skill shared by its author, an FPT PM, on 13 Oct 2026, adapted here with
the author's permission. Change: the requirement list may be any workbook with a 'Requirements' sheet
(or another sheet named by --sheet) whose header has an ID column and, optionally, Category, Priority,
Quote, Obligation and Status columns; the original expected the author's compliance-matrix builder.
Standard output is UTF-8, so a Windows console does not crash on symbols in slide text.

    python audit_proposal.py compliance-matrix.xlsx --deck deck.pptx [--docs a.docx b.md] \
        [-o audit-evidence.json]

WHAT THIS DOES: extracts the text of every deliverable location (one per slide for pptx --
title, body and speaker notes; one per heading section for docx/md/txt), then ranks those
locations against each requirement using idf-weighted term overlap, with a strong boost
when the requirement's own ID is written into the slide or its notes.

WHAT THIS DOES NOT DO: decide compliance. Term overlap tells you a topic was MENTIONED
somewhere. It cannot tell you the obligation was MET, and it cannot tell the difference
between "we provide 99.9% availability" and "the customer requires 99.9% availability,
which we do not offer". Every candidate below is a place for a human to go and read. The
script deliberately emits no status verdict -- you assign the status after reading the
evidence, and you record the locator you actually checked.

Two systematic blind spots to keep in mind:
  * A high score on a slide that merely restates the requirement is a false positive, and
    restating requirements is exactly what weak proposals do.
  * SUB and QUA requirements (forms, signatures, certificates, CVs) usually live outside
    the deck entirely, so a low score there is expected -- check the submission pack.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("error: openpyxl is required -- pip install openpyxl --break-system-packages")

# --------------------------------------------------------------------------------------
STOPWORDS = set("""
a an the and or but if then else of to in on at by for with without from into onto over under
is are was were be been being am do does did doing have has had having will would shall should
can could may might must ought this that these those it its it's as such than there their them
they he she we you your our us i not no nor only own same so too very just also both each few
more most other some any all such per via within across during before after above below between
than while when where which who whom whose what why how
system solution platform provide provided providing support supported bidder supplier vendor
contractor customer client requirement requirements required require shall must should ensure
ensuring able ability new proposed propose approach used use using including include includes
""".split())

WORD = re.compile(r"[A-Za-z][A-Za-z0-9_.+/-]*")
REQ_ID = re.compile(r"\b([A-Z]{2,4})[-_ ]?(\d{1,4})\b")


def normalise(token: str) -> str:
    t = token.lower().strip(".-/+_")
    if len(t) > 5 and t.endswith("ing"):
        t = t[:-3]
    elif len(t) > 5 and t.endswith("ed"):
        t = t[:-2]
    elif len(t) > 3 and t.endswith("s") and not t.endswith(("ss", "us", "is")):
        t = t[:-1]
    return t


def tokenise(text: str) -> list[str]:
    out = []
    for raw in WORD.findall(text or ""):
        if len(raw) < 2:
            continue
        t = normalise(raw)
        if len(t) < 3 or t in STOPWORDS or t.isdigit():
            continue
        out.append(t)
    return out


def canonical_ids(text: str) -> set[str]:
    """Every REQ-ID-shaped token in the text, normalised to PREFIX-NNN."""
    return {f"{m.group(1).upper()}-{int(m.group(2)):03d}" for m in REQ_ID.finditer(text or "")}


def canon(rid: str) -> str:
    m = REQ_ID.search((rid or "").upper())
    return f"{m.group(1)}-{int(m.group(2)):03d}" if m else (rid or "").upper()


# --------------------------------------------------------------------------------------
# extraction
# --------------------------------------------------------------------------------------
def extract_pptx(path: Path, skipped: list) -> list[dict]:
    try:
        from pptx import Presentation
    except ImportError:
        skipped.append(f"{path.name}: python-pptx not installed -- skipped")
        return []
    try:
        prs = Presentation(str(path))
    except Exception as exc:                                    # noqa: BLE001
        skipped.append(f"{path.name}: could not be opened ({exc}) -- skipped")
        return []

    def shape_text(shape) -> str:
        bits = []
        if shape.has_text_frame:
            bits.append(shape.text_frame.text)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                bits.append(" ".join(c.text for c in row.cells))
        if shape.shape_type == 6:                               # GROUP
            for sub in shape.shapes:
                bits.append(shape_text(sub))
        return "\n".join(b for b in bits if b)

    out = []
    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip()
        except Exception:                                       # noqa: BLE001
            pass
        body = "\n".join(filter(None, (shape_text(s) for s in slide.shapes)))
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text or ""
        if not (body.strip() or notes.strip()):
            continue
        out.append({"location": f"Slide {i}" + (f" — {title}" if title else ""),
                    "kind": "slide", "source": path.name, "index": i,
                    "title": title, "text": body, "notes": notes})
    return out


def extract_docx(path: Path, skipped: list) -> list[dict]:
    try:
        import docx
    except ImportError:
        skipped.append(f"{path.name}: python-docx not installed -- skipped "
                       "(pip install python-docx)")
        return []
    try:
        doc = docx.Document(str(path))
    except Exception as exc:                                    # noqa: BLE001
        skipped.append(f"{path.name}: could not be opened ({exc}) -- skipped")
        return []

    # Walk the body in document order so a table lands in the section it sits under --
    # doc.paragraphs and doc.tables are separate sequences and lose that interleaving.
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    sections, current = [], {"title": "(front matter)", "lines": []}
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = (para.text or "").strip()
            style = (para.style.name if para.style is not None else "") or ""
            if style.lower().startswith("heading") or style.lower() == "title":
                if current["lines"] or sections:
                    sections.append(current)
                current = {"title": text or "(untitled section)", "lines": []}
            elif text:
                current["lines"].append(text)
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                current["lines"].append(" ".join(c.text for c in row.cells))
    sections.append(current)

    return [{"location": f"{path.name} § {s['title']}", "kind": "section",
             "source": path.name, "index": i, "title": s["title"],
             "text": "\n".join(s["lines"]), "notes": ""}
            for i, s in enumerate(sections, start=1) if s["lines"]]


def extract_markdown(path: Path, skipped: list) -> list[dict]:
    try:
        raw = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        skipped.append(f"{path.name}: unreadable ({exc}) -- skipped")
        return []
    sections, current = [], {"title": "(front matter)", "lines": []}
    for line in raw.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            if current["lines"] or sections:
                sections.append(current)
            current = {"title": m.group(2).strip(), "lines": []}
        elif line.strip():
            current["lines"].append(line.strip())
    sections.append(current)
    return [{"location": f"{path.name} § {s['title']}", "kind": "section",
             "source": path.name, "index": i, "title": s["title"],
             "text": "\n".join(s["lines"]), "notes": ""}
            for i, s in enumerate(sections, start=1) if s["lines"]]


def extract(path: Path, skipped: list) -> list[dict]:
    if not path.exists():
        skipped.append(f"{path}: not found -- skipped")
        return []
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return extract_pptx(path, skipped)
    if suffix == ".docx":
        return extract_docx(path, skipped)
    if suffix in (".md", ".markdown", ".txt", ".rst"):
        return extract_markdown(path, skipped)
    skipped.append(f"{path.name}: no parser for '{suffix}' -- skipped "
                   f"(supported: .pptx .docx .md .txt .rst)")
    return []


# --------------------------------------------------------------------------------------
# requirements
# --------------------------------------------------------------------------------------
def read_requirements(path: Path, sheet: str = "Requirements") -> tuple[list[dict], str]:
    if path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        reqs = [{"id": r.get("id", ""), "category": (r.get("category") or "").upper(),
                 "priority": r.get("priority", ""), "obligation": r.get("obligation", ""),
                 "quote": r.get("quote", ""), "status": r.get("status", "")}
                for r in data.get("requirements", [])]
        return reqs, data.get("project", path.stem)

    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except FileNotFoundError:
        sys.exit(f"error: no such file: {path}")
    if sheet not in wb.sheetnames:
        sys.exit(f"error: {path} has no '{sheet}' sheet; name the requirement sheet with --sheet "
                 f"(sheets: {', '.join(wb.sheetnames)})")
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip().lower() for h in next(rows)]

    def col(*names, default=None):
        for n in names:
            for i, h in enumerate(header):
                if h.startswith(n):
                    return i
        return default

    ci = {"id": col("id"), "category": col("category"), "priority": col("priority"),
          "quote": col("quote"), "obligation": col("obligation"), "status": col("status")}
    if ci["id"] is None:
        sys.exit("error: the Requirements sheet has no 'ID' column")

    reqs = []
    for row in rows:
        rid = str(row[ci["id"]] or "").strip()
        if not rid:
            continue

        def get(key):
            i = ci[key]
            return str(row[i] or "").strip() if i is not None and i < len(row) else ""

        reqs.append({"id": rid, "category": get("category").upper(),
                     "priority": get("priority"), "obligation": get("obligation"),
                     "quote": get("quote"), "status": get("status")})
    wb.close()
    return reqs, (wb.properties.title or path.stem)


# --------------------------------------------------------------------------------------
# scoring
# --------------------------------------------------------------------------------------
def score_all(reqs, locations, top_n, threshold, id_boost):
    n_loc = len(locations)
    doc_tokens, doc_ids = [], []
    for loc in locations:
        toks = tokenise(loc["title"]) * 2 + tokenise(loc["text"]) + tokenise(loc["notes"])
        doc_tokens.append(Counter(toks))
        doc_ids.append(canonical_ids(loc["text"]) | canonical_ids(loc["notes"])
                       | canonical_ids(loc["title"]))

    df = Counter()
    for tf in doc_tokens:
        df.update(tf.keys())
    idf = {t: math.log(1 + n_loc / (1 + d)) for t, d in df.items()}
    default_idf = math.log(1 + n_loc)          # a term absent from every location

    results, unaddressed = [], []
    for r in reqs:
        # the obligation is the interpreted requirement, so it carries more signal than the
        # verbatim quote (which is full of the customer's legal boilerplate)
        qw = Counter()
        for t in tokenise(r["obligation"]):
            qw[t] += 2.0
        for t in tokenise(r["quote"]):
            qw[t] += 1.0
        mass = sum(w * idf.get(t, default_idf) for t, w in qw.items()) or 1.0
        rid = canon(r["id"])

        scored = []
        for i, loc in enumerate(locations):
            tf = doc_tokens[i]
            hit = {t: w * idf.get(t, default_idf) for t, w in qw.items() if tf.get(t)}
            overlap = sum(hit.values()) / mass
            tagged = rid in doc_ids[i]
            score = min(1.0, overlap + (id_boost if tagged else 0.0))
            if score <= 0:
                continue
            terms = sorted(hit, key=lambda t: -hit[t])[:6]
            scored.append({"location": loc["location"], "kind": loc["kind"],
                           "source": loc["source"], "score": round(score, 3),
                           "overlap": round(overlap, 3), "id_tagged": tagged,
                           "matched_terms": terms})
        scored.sort(key=lambda c: (-c["score"], c["location"]))
        best = scored[0]["score"] if scored else 0.0
        entry = {"id": r["id"], "category": r["category"], "priority": r["priority"],
                 "status_in_matrix": r["status"], "obligation": r["obligation"],
                 "best_score": round(best, 3), "id_tagged_anywhere": any(c["id_tagged"] for c in scored),
                 "candidates": scored[:top_n]}
        results.append(entry)
        if best < threshold:
            unaddressed.append({"id": r["id"], "category": r["category"],
                                "priority": r["priority"], "obligation": r["obligation"],
                                "best_score": round(best, 3),
                                "best_location": scored[0]["location"] if scored else None})
    return results, unaddressed


# --------------------------------------------------------------------------------------
def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(
        description="Match compliance-matrix requirements to candidate locations in the "
                    "proposal deliverables. Emits candidates for a human to verify -- never "
                    "a compliance verdict.")
    ap.add_argument("matrix", type=Path, help="compliance-matrix.xlsx (or requirements.json)")
    ap.add_argument("--sheet", default="Requirements",
                    help="sheet holding the requirement list (workbook input only)")
    ap.add_argument("--deck", type=Path, action="append", default=[],
                    help="proposal deck .pptx (repeatable)")
    ap.add_argument("--docs", type=Path, nargs="*", default=[],
                    help="additional deliverables: .docx .md .txt .rst")
    ap.add_argument("-o", "--output", type=Path, default=Path("audit-evidence.json"))
    ap.add_argument("--top", type=int, default=3, help="candidate locations per requirement")
    ap.add_argument("--threshold", type=float, default=0.22,
                    help="best score below this lands on likely_unaddressed (default 0.22)")
    ap.add_argument("--id-boost", type=float, default=0.55,
                    help="score added when the requirement ID is written on the location")
    args = ap.parse_args()

    reqs, project = read_requirements(args.matrix, args.sheet)
    if not reqs:
        sys.exit("error: no requirements found in the matrix")

    skipped: list[str] = []
    locations: list[dict] = []
    for p in list(args.deck) + list(args.docs):
        locations.extend(extract(Path(p), skipped))

    if not locations:
        print("\n  No readable deliverable content.", file=sys.stderr)
        for s in skipped:
            print(f"       - {s}", file=sys.stderr)
        sys.exit("\nerror: nothing to audit against -- pass --deck and/or --docs")

    results, unaddressed = score_all(reqs, locations, args.top, args.threshold, args.id_boost)

    payload = {
        "generated": datetime.now().isoformat(timespec="seconds"),
        "project": project,
        "matrix": str(args.matrix),
        "deliverables": sorted({loc["source"] for loc in locations}),
        "locations_indexed": len(locations),
        "skipped": skipped,
        "threshold": args.threshold,
        "method": "idf-weighted term overlap of obligation (x2) + verbatim quote (x1) against "
                  "per-location text, title (x2) and speaker notes, plus a "
                  f"+{args.id_boost} boost when the requirement ID appears in the location.",
        "caveat": "Scores indicate TOPIC MENTION, not compliance. A human must open each "
                  "candidate location, read it, and decide whether the obligation is actually "
                  "met before assigning a status. No verdict is implied by any score here.",
        "requirements": results,
        "likely_unaddressed": unaddressed,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")

    # ---- console summary --------------------------------------------------------------
    by_source = Counter(loc["source"] for loc in locations)
    tagged = sum(1 for r in results if r["id_tagged_anywhere"])
    print(f"\n  AUDIT EVIDENCE — {project}")
    print(f"  {'=' * 74}")
    print(f"  Indexed {len(locations)} locations "
          f"({', '.join(f'{v} in {k}' for k, v in by_source.items())})")
    print(f"  Scored  {len(reqs)} requirements; {tagged} of them appear as an explicit "
          f"ID tag in the content")
    for s in skipped:
        print(f"  SKIPPED {s}")

    print(f"\n  Requirements whose best candidate scores below {args.threshold} "
          f"— START HERE, these are the likely gaps:")
    if not unaddressed:
        print("       (none — every requirement has at least one plausible location)")
    order = {"Mandatory": 0, "Should": 1, "Optional": 2, "Informational": 3}
    for u in sorted(unaddressed, key=lambda u: (order.get(u["priority"], 9), u["id"])):
        flag = "!!" if u["priority"] == "Mandatory" else "  "
        note = " [SUB/QUA — likely in the submission pack, not the deck]" \
            if u["category"] in ("SUB", "QUA") else ""
        print(f"    {flag} {u['id']:<10} {u['priority']:<14} {u['best_score']:.2f}  "
              f"{u['obligation'][:52]}{note}")

    strong = [r for r in results if r["best_score"] >= args.threshold]
    print(f"\n  Top candidate per requirement ({len(strong)} with a plausible match):")
    for r in results:
        if not r["candidates"]:
            continue
        c = r["candidates"][0]
        tag = " [ID-tagged]" if c["id_tagged"] else ""
        mark = "?" if c["score"] < args.threshold else " "
        print(f"    {mark} {r['id']:<10} {c['score']:.2f}  {c['location'][:56]}{tag}")
        print(f"                  terms: {', '.join(c['matched_terms']) or '(none)'}")

    print(f"""
  {'-' * 74}
  READ THE CANDIDATES BEFORE YOU BELIEVE THEM. A score measures shared vocabulary,
  not compliance: a slide that merely repeats the customer's requirement scores just
  as highly as one that answers it. Open each location, decide whether the obligation
  is met, then record the status and the exact locator in the traceability matrix.

  -> {args.output}
""")


if __name__ == "__main__":
    main()
